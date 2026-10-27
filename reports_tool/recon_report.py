"""
Produce a sign-off style reconciliation workbook (Excel) for the migration.

    python -m reports_tool.recon_report            -> reports/migration_recon_report.xlsx
"""
from __future__ import annotations

import sqlite3
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from etl.migrate import run
from etl.seed_source import build

ROOT = Path(__file__).resolve().parent.parent
PAIRS = [("scp_item", "dim_product", "scp_item"), ("scp_location", "dim_site", None),
         ("scp_customer", "dim_customer", None), ("scp_forecast", "fact_demand_forecast", "scp_forecast"),
         ("scp_item_param", "planning_policy", "scp_item_param"), ("scp_user", "app_user", "scp_user")]
HEADER = PatternFill("solid", fgColor="1F4E79")
PASS = PatternFill("solid", fgColor="C6EFCE")
FAIL = PatternFill("solid", fgColor="FFC7CE")


def style_header(ws, row=1):
    for cell in ws[row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER
        cell.alignment = Alignment(horizontal="center")


def autosize(ws):
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col) + 2
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(width, 70)


def main(out: Path = ROOT / "reports" / "migration_recon_report.xlsx") -> Path:
    work = ROOT / "db"
    source = build(work / "scp_classic.db")
    target = run(source, work / "xpr_nextgen.db")
    con = sqlite3.connect(target)
    con.execute("ATTACH DATABASE ? AS src", (str(source),))
    one = lambda sql, p=(): con.execute(sql, p).fetchone()[0]

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Source table", "Target table", "Source rows", "Target rows", "Rejected", "Difference", "Status"])
    style_header(ws)
    for s, t, r in PAIRS:
        src, tgt = one(f"SELECT COUNT(*) FROM src.{s}"), one(f"SELECT COUNT(*) FROM {t}")
        rej = one("SELECT COUNT(*) FROM migration_reject WHERE source_table=?", (r,)) if r else 0
        diff = src - tgt - rej
        ws.append([s, t, src, tgt, rej, diff, "PASS" if diff == 0 else "FAIL"])
        ws.cell(ws.max_row, 7).fill = PASS if diff == 0 else FAIL
    ws.append([])
    ws.append(["Generated", datetime.now().strftime("%Y-%m-%d %H:%M")])
    autosize(ws)

    ws = wb.create_sheet("Validation SQL")
    ws.append(["Check", "Offending rows", "Status"])
    style_header(ws)
    for f in sorted((ROOT / "sql" / "validation").glob("*.sql")):
        n = len(con.execute(f.read_text()).fetchall())
        ws.append([f.stem, n, "PASS" if n == 0 else "FAIL"])
        ws.cell(ws.max_row, 3).fill = PASS if n == 0 else FAIL
    autosize(ws)

    ws = wb.create_sheet("Reject summary")
    ws.append(["Source table", "Reason code", "Rows"])
    style_header(ws)
    for row in con.execute("SELECT source_table, reason_code, COUNT(*) FROM migration_reject GROUP BY 1, 2 ORDER BY 1, 3 DESC"):
        ws.append(list(row))
    autosize(ws)

    ws = wb.create_sheet("Reject detail")
    ws.append(["Reject ID", "Source table", "Source key", "Reason code", "Detail"])
    style_header(ws)
    for row in con.execute("SELECT * FROM migration_reject WHERE reason_code <> 'PARENT_REJECTED' ORDER BY reject_id"):
        ws.append(list(row))
    autosize(ws)

    ws = wb.create_sheet("Qty by product")
    ws.append(["Product", "Source qty (each)", "Target qty (each)", "Status"])
    style_header(ws)
    for pid, tgt in con.execute("SELECT product_id, SUM(qty_each) FROM fact_demand_forecast GROUP BY 1 ORDER BY 1").fetchall():
        src = one("""SELECT SUM(CAST(ROUND(s.fc_qty * CASE WHEN i.uom='CS' THEN i.case_pack ELSE 1 END + 1e-9) AS INTEGER))
                     FROM src.scp_forecast s JOIN src.scp_item i ON i.item_code = s.item_code
                     WHERE s.fc_id IN (SELECT source_fc_id FROM fact_demand_forecast WHERE product_id = ?)""", (pid,))
        ws.append([pid, src, tgt, "PASS" if src == tgt else "FAIL"])
        ws.cell(ws.max_row, 4).fill = PASS if src == tgt else FAIL
    autosize(ws)

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


if __name__ == "__main__":
    print("Report written to", main())
